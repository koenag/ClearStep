"""Deterministic parsing. No files or excerpts are sent to a model here."""

import csv
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import DomainError
from app.services.ingestion import aggregate, parse_csv

SUPPORTED = {".csv", ".xlsx", ".txt", ".pdf"}


def parse_file(filename, content, settings, sheet=None):
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise DomainError("unsupported_format", "Supported: CSV, XLSX, TXT, text-based PDF.", 415)
    if len(content) > settings.max_upload_bytes:
        raise DomainError("upload_too_large", "Upload exceeds the configured limit.", 413)
    if not content:
        raise DomainError("empty_file", "The file is empty.")
    try:
        if suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > settings.max_upload_bytes * 20:
                    raise DomainError(
                        "expanded_file_too_large", "Expanded workbook exceeds the limit."
                    )
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                names = book.sheetnames
                if sheet is None and len(names) != 1:
                    raise DomainError(
                        "sheet_required", "Select a sheet explicitly: " + ", ".join(names)
                    )
                chosen = sheet or names[0]
                if chosen not in names:
                    raise DomainError("invalid_sheet", "Selected sheet does not exist.")
                rows = []
                for row in book[chosen].iter_rows():
                    if len(rows) > settings.max_rows:
                        raise DomainError("too_many_rows", "Workbook exceeds row limit.")
                    if len(row) > 100:
                        raise DomainError("too_many_columns", "At most 100 columns are supported.")
                    if any(c.data_type == "f" for c in row):
                        raise DomainError(
                            "formula_review_required",
                            "Export calculated values: spreadsheet formulas are not evaluated.",
                        )
                    values = [
                        c.value.isoformat() if isinstance(c.value, (date, datetime)) else c.value
                        for c in row
                    ]
                    if any(v is not None for v in values):
                        rows.append(values)
                if not rows:
                    raise DomainError("empty_sheet", "Selected sheet is empty.")
                frame = pd.DataFrame(rows[1:], columns=[str(v or "") for v in rows[0]])
                content = frame.to_csv(index=False).encode()
                location = f"sheet {chosen}"
            finally:
                book.close()
        else:
            location = "CSV"
        if suffix in {".csv", ".xlsx"}:
            # Bound dimensions before row-by-row financial validation.
            frame = pd.read_csv(io.BytesIO(content), dtype=str, keep_default_na=False)
            if len(frame) > settings.max_rows or len(frame.columns) > 100:
                raise DomainError("table_too_large", "Table exceeds row/column limits.")
            headers = next(csv.reader(io.StringIO(content.decode("utf-8-sig"))), [])
            normalized = [h.strip().lower().replace(" ", "_") for h in headers]
            if len(set(normalized)) != len(normalized):
                raise DomainError("duplicate_columns", "Duplicate normalized columns.")
            columns = set(normalized)
            # Legacy cost-enriched sales tables remain supported. Ordinary POS
            # exports do not contain COGS; never substitute zero for missing costs.
            from app.services.ingestion import ALIASES

            canonical = {ALIASES.get(h, h) for h in normalized}
            if {"date", "product", "quantity", "unit_price", "unit_cost"} <= canonical:
                records, warnings = parse_csv(content)
                baseline = aggregate(records)
                count = len(records)
            elif {
                "date",
                "receipt",
                "item",
                "qty",
                "gross_sales",
                "discount",
                "net_sales",
                "tax",
            } <= columns:
                frame.columns = normalized
                baseline, warnings = aggregate_pos(frame)
                count = baseline["line_items"]
            else:
                baseline = None
                count = len(frame)
                warnings = [
                    "Generic table retained as source evidence; no financial baseline inferred from these columns."
                ]
            return {
                "format": suffix[1:],
                "chunks": [{"location": location, "text": content.decode("utf-8-sig")}],
                "baseline": baseline,
                "warnings": warnings,
                "rows": count,
                "columns": list(frame.columns),
            }
        if suffix == ".txt":
            text = content.decode("utf-8-sig")
            chunks = [
                {
                    "location": f"lines {i + 1}-{min(i + 60, len(text.splitlines()))}",
                    "text": "\n".join(text.splitlines()[i : i + 60]),
                }
                for i in range(0, len(text.splitlines()), 60)
            ]
        else:
            pdf = PdfReader(io.BytesIO(content))
            if pdf.is_encrypted:
                raise DomainError("encrypted_pdf", "Encrypted PDFs are unsupported.")
            if len(pdf.pages) > settings.max_pdf_pages:
                raise DomainError("too_many_pages", "PDF exceeds page limit.")
            chunks = [
                {"location": f"page {i + 1}", "text": page.extract_text() or ""}
                for i, page in enumerate(pdf.pages)
            ]
        length = sum(len(c["text"]) for c in chunks)
        if length > settings.max_document_chars:
            raise DomainError("document_too_large", "Extracted document exceeds text limit.")
        if not any(c["text"].strip() for c in chunks):
            raise DomainError(
                "no_extractable_text",
                "No usable text found. Scanned documents require OCR, which is unsupported.",
            )
        return {
            "format": suffix[1:],
            "chunks": chunks,
            "baseline": None,
            "warnings": [
                "Extracted numerical fields require explicit review before financial use."
            ],
            "rows": None,
            "columns": [],
        }
    except DomainError:
        raise
    except (ValueError, UnicodeError, zipfile.BadZipFile) as exc:
        raise DomainError("parse_failed", "Cannot parse file: " + str(exc)[:300]) from exc
    except Exception as exc:
        raise DomainError("parse_failed", "The document is invalid or unsupported.") from exc


def aggregate_pos(frame):
    """Observed POS revenue excluding tax, with explicitly unavailable COGS/margin."""
    from app.financial.engine import money

    revenue = gross_total = discount_total = tax_total = units = Decimal(0)
    products, monthly, transactions, dates, invalid = {}, {}, set(), [], []
    for i, row in frame.iterrows():
        try:
            day = date.fromisoformat(row["date"].strip()).isoformat()
            quantity, gross, discount, net, tax = [
                Decimal(row[k]) for k in ["qty", "gross_sales", "discount", "net_sales", "tax"]
            ]
            if (
                any(
                    not n.is_finite() or n < 0 or n > 10000000
                    for n in [quantity, gross, discount, net, tax]
                )
                or quantity == 0
            ):
                raise ValueError()
            if gross - discount != net or not row["receipt"].strip() or not row["item"].strip():
                raise ValueError()
            receipt, item = row["receipt"].strip(), row["item"].strip()
        except (ValueError, ArithmeticError):
            invalid.append(int(i) + 1)
            continue
        revenue += net
        gross_total += gross
        discount_total += discount
        tax_total += tax
        units += quantity
        transactions.add((day, receipt))
        dates.append(day)
        product = products.setdefault(
            item, {"product": item, "revenue": Decimal(0), "units": Decimal(0), "cost": None}
        )
        product["revenue"] += net
        product["units"] += quantity
        monthly[day[:7]] = monthly.get(day[:7], Decimal(0)) + net
    if not dates:
        raise DomainError(
            "no_valid_pos_rows",
            "No valid POS rows with ISO dates and reconciled net sales were found.",
        )
    warnings = (
        [f"Skipped invalid POS data rows (excluding header): {invalid[:20]}"] if invalid else []
    )
    result = {
        "revenue": money(revenue),
        "gross_sales": money(gross_total),
        "discounts": money(discount_total),
        "tax_collected": money(tax_total),
        "variable_costs": None,
        "contribution_profit": None,
        "contribution_margin": None,
        "units": float(units),
        "transactions": len(transactions),
        "average_transaction_value": money(revenue / len(transactions)),
        "line_items": len(dates),
        "period_start": min(dates),
        "period_end": max(dates),
        "products": [
            {**p, "revenue": money(p["revenue"]), "units": float(p["units"])}
            for p in products.values()
        ],
        "monthly": [{"month": m, "revenue": money(v)} for m, v in sorted(monthly.items())],
        "warnings": [
            "Revenue is recorded net sales after discounts, excluding collected tax.",
            "POS records contain no product costs. Variable costs, contribution profit/margin, and net profit are unavailable.",
            "Transaction count uses distinct date + receipt pairs; sales lines are not unique checkouts.",
            "Units are menu items, not pounds of coffee. Recorded periods are not extrapolated.",
        ],
    }
    return result, warnings
